"""
Generador de registros de capacitación en Excel.

Plantilla: data/plantillas/plantilla_asistencia.xlsx
Salida:    documentos_generados/Asistencia_YYYYMMDD_HHMMSS.xlsx
"""
import shutil
from datetime import datetime

import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment

from src.config import (
    CARPETA_FIRMAS,
    CARPETA_SALIDA,
    PLANTILLA_ASISTENCIA,
)


def guardar_registro_capacitacion(
    tema: str,
    fecha: str,
    hora: str,
    trabajadores: list[dict],
) -> str:
    """
    Genera el Excel de asistencia a capacitación.

    trabajadores: lista de dicts con claves 'dni', 'nombre', 'cargo'
    Retorna la ruta del archivo generado.
    """
    if not PLANTILLA_ASISTENCIA.exists():
        raise FileNotFoundError(f"Plantilla no encontrada: {PLANTILLA_ASISTENCIA}")

    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)

    nombre_archivo = f"Asistencia_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    ruta_destino   = CARPETA_SALIDA / nombre_archivo
    shutil.copyfile(str(PLANTILLA_ASISTENCIA), str(ruta_destino))

    wb = openpyxl.load_workbook(str(ruta_destino))
    ws = wb['Lista Asistencia']

    ws['J14']  = tema
    ws['AI14'] = fecha
    ws['AI15'] = hora

    for i, t in enumerate(trabajadores):
        fila     = 18 + i
        dni      = t['dni']
        apellido = (t.get('apellido') or '').strip()
        nombre_p = (t.get('nombre')   or '').strip()
        nombre   = f"{apellido} {nombre_p}".strip() if apellido else nombre_p
        cargo    = t.get('cargo') or ''

        ws.merge_cells(f'B{fila}:M{fila}')
        ws[f'B{fila}'] = nombre
        ws[f'N{fila}'] = dni
        ws[f'R{fila}'] = 'TALLER/PLANTA'

        cell_cargo = ws[f'X{fila}']
        cell_cargo.value = cargo
        cell_cargo.alignment = Alignment(
            wrap_text=True, horizontal='center', vertical='center'
        )

        img_path = CARPETA_FIRMAS / f"firma_{dni}.png"
        if img_path.exists():
            img = ExcelImage(str(img_path))
            img.anchor = f'AI{fila}'
            ws.add_image(img)

    wb.save(str(ruta_destino))
    return str(ruta_destino)
