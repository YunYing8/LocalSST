"""
Generador de Cargo de Entrega de EPP en Excel.

Carga la plantilla EPPs.xlsx y rellena los datos del trabajador y EPPs.
Salida: documentos_generados/EPP_{DNI}_{YYYYMMDD_HHMMSS}.xlsx
"""
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

from src.config import CARPETA_FIRMAS, CARPETA_SALIDA, FIRMA_ALTO_PX, PROJECT_ROOT

PLANTILLA_EPP = PROJECT_ROOT / "data" / "plantillas" / "EPPs.xlsx"

EPPS_BASICOS = [
    "Casco de seguridad",
    "Lentes de seguridad",
    "Protector auditivo",
    "Mascarilla / Respirador",
    "Guantes de seguridad",
    "Zapatos de seguridad",
    "Chaleco reflectivo",
]


def generar_registro_epp(trabajador: dict, epps: list[str], fecha: str) -> str:
    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)

    apellido = (trabajador.get("apellido") or "").strip()
    nombre_p = (trabajador.get("nombre")   or "").strip()
    nombre   = f"{apellido} {nombre_p}".strip() if apellido else nombre_p
    dni      = trabajador.get("dni", "")
    cargo    = trabajador.get("cargo") or ""

    wb = load_workbook(str(PLANTILLA_EPP))
    ws = wb.active

    # ── Fecha (celda info superior derecha) ──────────────────────────────────
    ws["I3"] = f"Fecha: {fecha}"

    # ── Datos del trabajador ─────────────────────────────────────────────────
    ws["C12"] = nombre
    ws["J12"] = dni
    ws["G13"] = cargo

    # ── Filas de EPP (desde fila 16) ─────────────────────────────────────────
    firma_path = CARPETA_FIRMAS / f"firma_{dni}.png"
    row_h      = max(FIRMA_ALTO_PX, 50) * 0.75 + 8

    for i, epp in enumerate(epps):
        fila = 16 + i
        ws.row_dimensions[fila].height = row_h

        ws[f"A{fila}"] = i + 1
        ws[f"B{fila}"] = fecha
        ws[f"C{fila}"] = epp

        if firma_path.exists():
            img        = ExcelImage(str(firma_path))
            img.height = FIRMA_ALTO_PX
            img.width  = int(FIRMA_ALTO_PX * 2.5)
            img.anchor = f"E{fila}"
            ws.add_image(img)

    ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"EPP_{dni}_{ts}.xlsx"
    ruta     = CARPETA_SALIDA / filename
    wb.save(str(ruta))
    return str(ruta)
