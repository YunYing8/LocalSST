"""
Generador de Cargo de Entrega de EPP en Excel.

Carga la plantilla EPPs.xlsx y rellena solo las celdas dinámicas.
Salida: documentos_generados/EPP_{DNI}_{YYYYMMDD_HHMMSS}.xlsx
"""
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

from src.config import CARPETA_FIRMAS, CARPETA_SALIDA, FIRMA_ALTO_PX, PLANTILLA_EPPS

# Filas de la tabla de EPPs en la plantilla (filas 16 a 21 = 6 entregas)
_FILA_INICIO_EPPS = 16
_MAX_EPPS         = 6


def generar_registro_epp(trabajador: dict, epps: list[str], fecha: str,
                         area: str = "Taller", nro_registro: str = "") -> str:
    """
    Genera el Excel de Cargo de Entrega de EPP para un trabajador.

    trabajador   : dict con claves dni, nombre, apellido, cargo
    epps         : lista de nombres de equipos entregados (máx 6)
    fecha        : fecha de entrega (dd/mm/yyyy)
    area         : área del trabajador (por defecto "Taller")
    nro_registro : número de registro (opcional)
    Retorna      : ruta del archivo generado
    """
    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)

    apellido = (trabajador.get("apellido") or "").strip()
    nombre_p = (trabajador.get("nombre")   or "").strip()
    nombre   = f"{apellido} {nombre_p}".strip() if apellido else nombre_p
    dni      = trabajador.get("dni", "")
    cargo    = trabajador.get("cargo") or ""

    # Cargar plantilla sin modificar el original
    wb = load_workbook(str(PLANTILLA_EPPS))
    ws = wb["ENTREGA"]

    # ── Celdas dinámicas ──────────────────────────────────────────────────────
    ws["C5"]  = nro_registro   # N° Registro
    ws["C12"] = nombre         # Nombre del trabajador
    ws["J12"] = dni            # DNI
    ws["C13"] = area           # Área (por defecto "Taller")
    ws["H13"] = cargo          # Puesto / Cargo

    # ── Tabla de EPPs (filas 16–21) ───────────────────────────────────────────
    firma_path = CARPETA_FIRMAS / f"firma_{dni}.png"

    for i in range(_MAX_EPPS):
        fila = _FILA_INICIO_EPPS + i
        epp  = epps[i] if i < len(epps) else ""

        ws[f"B{fila}"] = fecha if epp else ""   # Fecha
        ws[f"C{fila}"] = epp                    # Equipo entregado

        # Firma del trabajador (RECIBÍ) como imagen en columna F
        if epp and firma_path.exists():
            img        = ExcelImage(str(firma_path))
            img.height = FIRMA_ALTO_PX
            img.width  = int(FIRMA_ALTO_PX * 2.5)
            img.anchor = f"F{fila}"
            ws.add_image(img)

    # ── Guardar como nuevo archivo ────────────────────────────────────────────
    ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"EPP_{dni}_{ts}.xlsx"
    ruta     = CARPETA_SALIDA / filename
    wb.save(str(ruta))
    return str(ruta)
