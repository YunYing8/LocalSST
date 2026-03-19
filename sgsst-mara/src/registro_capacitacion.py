"""
Generador de registros de capacitación en Excel.

Plantilla esperada:
    data/plantillas/plantilla_asistencia.xlsx

Salida:
    documentos_generados/Asistencia_YYYYMMDD_HHMMSS.xlsx
"""
import os
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment

from src.database import DB_PATH

PROJECT_ROOT     = DB_PATH.parent.parent
RUTA_PLANTILLA   = PROJECT_ROOT / "data" / "plantillas" / "plantilla_asistencia.xlsx"
CARPETA_FIRMAS   = PROJECT_ROOT / "data" / "firmas"
CARPETA_SALIDA   = PROJECT_ROOT / "documentos_generados"


def guardar_registro_capacitacion(
    tema: str,
    fecha: str,
    hora: str,
    trabajadores: list[dict],
) -> str:
    """
    Genera el Excel de asistencia a capacitación.

    trabajadores: lista de dicts con claves 'dni', 'nombre', 'cargo'
    Retorna la ruta del archivo generado.
    """
    if not RUTA_PLANTILLA.exists():
        raise FileNotFoundError(f"Plantilla no encontrada: {RUTA_PLANTILLA}")

    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)

    nombre_archivo = f"Asistencia_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    ruta_destino   = CARPETA_SALIDA / nombre_archivo
    shutil.copyfile(str(RUTA_PLANTILLA), str(ruta_destino))

    wb = openpyxl.load_workbook(str(ruta_destino))
    ws = wb['Lista Asistencia']

    ws['J14'] = tema
    ws['AI14'] = fecha
    ws['AI15'] = hora

    fila = 18
    for t in trabajadores:
        dni    = t['dni']
        nombre = t['nombre']
        cargo  = t.get('cargo') or ''

        ws.merge_cells(f'B{fila}:M{fila}')
        ws[f'B{fila}'] = nombre
        ws[f'N{fila}'] = dni
        ws[f'R{fila}'] = 'TALLER/PLANTA'

        cell_cargo = ws[f'X{fila}']
        cell_cargo.value = cargo
        cell_cargo.alignment = Alignment(
            wrap_text=True, horizontal='center', vertical='center'
        )

        img_path = CARPETA_FIRMAS / f"firma_{dni}.png"
        if img_path.exists():
            img = ExcelImage(str(img_path))
            img.anchor = f'AI{fila}'
            ws.add_image(img)

        fila += 1

    wb.save(str(ruta_destino))
    return str(ruta_destino)
