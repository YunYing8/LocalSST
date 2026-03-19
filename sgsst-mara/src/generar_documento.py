from pathlib import Path
from datetime import datetime

from docxtpl import DocxTemplate, InlineImage
from docx.shared import Cm, Pt
from docx2pdf import convert

from src.database import get_connection, DB_PATH

PROJECT_ROOT   = DB_PATH.parent.parent
RUTA_PLANTILLA = PROJECT_ROOT / "data" / "plantillas" / "plantilla_risst.docx"
CARPETA_FIRMAS = PROJECT_ROOT / "data" / "firmas"
CARPETA_SALIDA = PROJECT_ROOT / "documentos_generados"

_MESES = {
    1: 'enero',    2: 'febrero',  3: 'marzo',    4: 'abril',
    5: 'mayo',     6: 'junio',    7: 'julio',     8: 'agosto',
    9: 'septiembre', 10: 'octubre', 11: 'noviembre', 12: 'diciembre',
}


def convertir_fecha(fecha_str: str) -> str:
    """'dd/mm/yyyy'  →  'Lima, D de mes del YYYY'"""
    try:
        fecha = datetime.strptime(fecha_str.strip(), '%d/%m/%Y')
        return f"Lima, {fecha.day} de {_MESES[fecha.month]} del {fecha.year}"
    except Exception:
        return fecha_str


def _aplicar_fuente(doc: DocxTemplate) -> None:
    """Aplica Century Gothic 11pt a todos los runs del documento."""
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Century Gothic'
            run.font.size = Pt(11)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Century Gothic'
                        run.font.size = Pt(11)


def generar_constancia(
    dni: str,
    nombre: str,
    apellido: str,
    cargo: str,
    fecha: str,
) -> tuple[bool, str]:
    """
    Genera la constancia RISST para un trabajador y la convierte a PDF.
    Retorna (True, nombre_pdf) o (False, mensaje_error).
    """
    try:
        CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)

        if not RUTA_PLANTILLA.exists():
            return False, f"Plantilla no encontrada: {RUTA_PLANTILLA}"

        doc = DocxTemplate(str(RUTA_PLANTILLA))

        ruta_firma = CARPETA_FIRMAS / f"firma_{dni}.png"

        contexto = {
            'NOMBRE':    nombre,
            'APELLIDOS': apellido or nombre,
            'DNI':       dni,
            'CARGO':     cargo or '',
            'FECHA':     convertir_fecha(fecha),
        }
        if ruta_firma.exists():
            contexto['FIRMA'] = InlineImage(doc, str(ruta_firma), width=Cm(4))

        doc.render(contexto)
        _aplicar_fuente(doc)

        nombre_docx = f"constancia_RISST_{nombre.replace(' ', '_')}.docx"
        ruta_docx   = CARPETA_SALIDA / nombre_docx
        doc.save(str(ruta_docx))

        convert(str(ruta_docx))
        nombre_pdf = nombre_docx.replace('.docx', '.pdf')

        # Registrar en historial
        with get_connection() as conn:
            conn.execute(
                """INSERT INTO documentos_generados
                   (dni_trabajador, tipo_documento, fecha_generacion, ruta_archivo)
                   VALUES (?, ?, ?, ?)""",
                (
                    dni,
                    'CONSTANCIA_RISST',
                    datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                    str(CARPETA_SALIDA / nombre_pdf),
                ),
            )

        return True, nombre_pdf

    except Exception as exc:
        return False, str(exc)


def generar_todas_constancias(fecha: str) -> dict:
    """
    Genera constancias para todos los trabajadores ACTIVOS.
    Retorna {"total", "exitosos", "fallidos", "log"}.
    """
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT dni, nombre, apellido, cargo FROM trabajadores "
            "WHERE estado = 'ACTIVO' ORDER BY nombre ASC"
        ).fetchall()
    trabajadores = [dict(r) for r in rows]

    exitosos = 0
    fallidos = 0
    log: list[str] = []

    for t in trabajadores:
        ok, msg = generar_constancia(
            t['dni'],
            t['nombre'],
            t['apellido'] or '',
            t['cargo'] or '',
            fecha,
        )
        if ok:
            exitosos += 1
            log.append(f"OK   {t['nombre']}  →  {msg}")
        else:
            fallidos += 1
            log.append(f"ERR  {t['nombre']}  →  {msg}")

    return {
        "total":    len(trabajadores),
        "exitosos": exitosos,
        "fallidos": fallidos,
        "log":      log,
    }
