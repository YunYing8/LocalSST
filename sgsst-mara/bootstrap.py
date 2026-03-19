"""
Ejecutar UNA SOLA VEZ desde la raiz del proyecto:
    python bootstrap.py
Crea todos los archivos fuente del sistema SGSST-MARA.
"""
from pathlib import Path

BASE = Path(__file__).parent

files = {}

# ── requirements.txt ────────────────────────────────────────────────────────
files["requirements.txt"] = """\
python-docx==1.1.2
openpyxl==3.1.2
Pillow==10.3.0
docxtpl==0.19.0
docx2pdf==0.1.8
requests==2.32.3
rembg==2.0.57
"""

# ── src/__init__.py ──────────────────────────────────────────────────────────
files["src/__init__.py"] = ""

# ── src/database.py ──────────────────────────────────────────────────────────
files["src/database.py"] = '''\
import sqlite3
from pathlib import Path

DB_PATH = Path(__file__).resolve().parent.parent / "data" / "sgsst.db"

_SCHEMA = """
CREATE TABLE IF NOT EXISTS trabajadores (
    id               INTEGER PRIMARY KEY AUTOINCREMENT,
    dni              TEXT UNIQUE NOT NULL,
    nombre           TEXT NOT NULL,
    apellido         TEXT,
    cargo            TEXT,
    fecha_nacimiento TEXT,
    correo           TEXT,
    celular          TEXT,
    estado           TEXT NOT NULL DEFAULT \'ACTIVO\',
    created_at       TEXT
);

CREATE TABLE IF NOT EXISTS documentos_generados (
    id               INTEGER PRIMARY KEY AUTOINCREMENT,
    dni_trabajador   TEXT NOT NULL,
    tipo_documento   TEXT NOT NULL,
    fecha_generacion TEXT,
    ruta_archivo     TEXT NOT NULL,
    FOREIGN KEY (dni_trabajador) REFERENCES trabajadores(dni)
);
"""


def get_connection() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with get_connection() as conn:
        conn.executescript(_SCHEMA)
        try:
            conn.execute("ALTER TABLE trabajadores ADD COLUMN apellido TEXT")
            conn.commit()
        except Exception:
            pass
'''

# ── src/trabajadores.py ──────────────────────────────────────────────────────
files["src/trabajadores.py"] = '''\
from datetime import datetime
from src.database import get_connection

_ESTADOS_VALIDOS = {"ACTIVO", "INACTIVO"}


def buscar_por_dni(dni: str) -> dict | None:
    with get_connection() as conn:
        row = conn.execute(
            "SELECT * FROM trabajadores WHERE dni = ?", (dni.strip(),)
        ).fetchone()
    return dict(row) if row else None


def listar_activos() -> list[dict]:
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT * FROM trabajadores WHERE estado = \'ACTIVO\' ORDER BY nombre ASC"
        ).fetchall()
    return [dict(r) for r in rows]


def cambiar_estado(dni: str, nuevo_estado: str) -> bool:
    if nuevo_estado not in _ESTADOS_VALIDOS:
        raise ValueError(f"Estado invalido \'{nuevo_estado}\'. Validos: {_ESTADOS_VALIDOS}")
    with get_connection() as conn:
        cursor = conn.execute(
            "UPDATE trabajadores SET estado = ? WHERE dni = ?", (nuevo_estado, dni.strip())
        )
    return cursor.rowcount > 0


def registrar_trabajador(datos: dict) -> bool:
    created_at = datetime.now().strftime(\'%Y-%m-%d %H:%M:%S\')
    with get_connection() as conn:
        cursor = conn.execute(
            """INSERT OR IGNORE INTO trabajadores
               (dni, nombre, apellido, cargo, fecha_nacimiento, correo, celular, estado, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                datos["dni"], datos["nombre"], datos.get("apellido"),
                datos.get("cargo"), datos.get("fecha_nacimiento"),
                datos.get("correo"), datos.get("celular"),
                datos.get("estado", "ACTIVO"), created_at,
            ),
        )
    return cursor.rowcount > 0
'''

# ── src/importar_excel.py ────────────────────────────────────────────────────
files["src/importar_excel.py"] = '''\
import openpyxl
from datetime import datetime
from src.database import get_connection

_INSERT = """
INSERT OR IGNORE INTO trabajadores
    (dni, nombre, apellido, cargo, fecha_nacimiento, correo, celular, estado, created_at)
VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
"""


def importar_trabajadores(ruta_excel: str) -> dict:
    wb = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    importados = omitidos = 0
    errores = []
    conn = get_connection()
    try:
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                nombre_col, apellido_col, dni_raw, cargo, fecha_nac, _edad, correo, celular, estado = row
                dni = str(dni_raw).strip() if dni_raw is not None else None
                if not dni or dni == "None":
                    omitidos += 1
                    continue
                nombre_raw = str(nombre_col).strip() if nombre_col else ""
                if not nombre_raw or nombre_raw == "None":
                    omitidos += 1
                    continue
                apellido_str = str(apellido_col).strip() if apellido_col else ""
                nombre       = f"{apellido_str} {nombre_raw}".strip()
                estado_val   = str(estado).strip().upper() if estado else "ACTIVO"
                fecha_str    = str(fecha_nac).strip() if fecha_nac is not None else None
                cargo_str    = str(cargo).strip() if cargo is not None else None
                correo_str   = str(correo).strip() if correo is not None else None
                celular_str  = str(celular).strip() if celular is not None else None
                created_at   = datetime.now().strftime(\'%Y-%m-%d %H:%M:%S\')
                cursor = conn.execute(
                    _INSERT,
                    (dni, nombre, apellido_str, cargo_str, fecha_str,
                     correo_str, celular_str, estado_val, created_at),
                )
                if cursor.rowcount == 1:
                    importados += 1
                else:
                    omitidos += 1
            except Exception as exc:
                errores.append(f"Fila {row_idx}: {exc}")
                omitidos += 1
        conn.commit()
    finally:
        conn.close()
        wb.close()
    return {"importados": importados, "omitidos": omitidos, "errores": errores}
'''

# ── src/importar_sheets.py ───────────────────────────────────────────────────
files["src/importar_sheets.py"] = '''\
import csv
import requests
from datetime import datetime
from src.database import get_connection

GOOGLE_SHEET_ID = \'1N0i-sG3UBH2UMVo_DF4Ivcwxyr2MvgIfAa4h18oQ5oc\'

_UPSERT = """
INSERT INTO trabajadores
    (dni, nombre, apellido, cargo, fecha_nacimiento, correo, celular, estado, created_at)
VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
ON CONFLICT(dni) DO UPDATE SET
    nombre=excluded.nombre, apellido=excluded.apellido,
    cargo=excluded.cargo, fecha_nacimiento=excluded.fecha_nacimiento,
    correo=excluded.correo, celular=excluded.celular, estado=excluded.estado
"""


def sincronizar_desde_sheets() -> dict:
    url = (f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}"
           f"/gviz/tq?tqx=out:csv&sheet=Registro")
    response = requests.get(url, timeout=15)
    response.raise_for_status()
    lineas = list(csv.reader(response.content.decode(\'utf-8\').splitlines()))
    if not lineas:
        return {"nuevos": 0, "actualizados": 0, "omitidos": 0, "errores": ["Hoja vacia"]}
    enc = lineas[0]
    def _idx(c): return enc.index(c) if c in enc else None
    idx_nombre=enc.index(\'Nombre\'); idx_apellido=enc.index(\'Apellido\')
    idx_dni=enc.index(\'DNI\'); idx_cargo=enc.index(\'Cargo\'); idx_estado=enc.index(\'Estado\')
    idx_fn=_idx(\'Fecha de nacimiento\'); idx_correo=_idx(\'Correo electronico\'); idx_celular=_idx(\'Celular\')
    nuevos=actualizados=omitidos=0; errores=[]
    conn = get_connection()
    try:
        for i, fila in enumerate(lineas[1:], start=2):
            try:
                if len(fila) <= idx_estado: omitidos+=1; continue
                dni = str(fila[idx_dni]).strip()
                if not dni: omitidos+=1; continue
                apellido=fila[idx_apellido].strip(); nombre_p=fila[idx_nombre].strip()
                if not apellido and not nombre_p: omitidos+=1; continue
                nombre_completo=f"{apellido} {nombre_p}".strip()
                cargo   = fila[idx_cargo].strip() if len(fila)>idx_cargo else None
                estado  = fila[idx_estado].strip().upper() if len(fila)>idx_estado else \'ACTIVO\'
                fecha_nac = fila[idx_fn].strip() if idx_fn and len(fila)>idx_fn else None
                correo    = fila[idx_correo].strip() if idx_correo and len(fila)>idx_correo else None
                celular   = fila[idx_celular].strip() if idx_celular and len(fila)>idx_celular else None
                created_at= datetime.now().strftime(\'%Y-%m-%d %H:%M:%S\')
                existe = conn.execute("SELECT 1 FROM trabajadores WHERE dni=?", (dni,)).fetchone()
                conn.execute(_UPSERT,(dni,nombre_completo,apellido,cargo,fecha_nac,correo,celular,estado,created_at))
                actualizados+=1 if existe else None; nuevos+= 0 if existe else 1
                if existe: actualizados+=1
                else: nuevos+=1
            except Exception as exc:
                errores.append(f"Fila {i}: {exc}")
        conn.commit()
    finally:
        conn.close()
    return {"nuevos": nuevos, "actualizados": actualizados, "omitidos": omitidos, "errores": errores}
'''

# ── src/generar_documento.py ─────────────────────────────────────────────────
files["src/generar_documento.py"] = '''\
from pathlib import Path
from datetime import datetime
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Cm, Pt
from docx2pdf import convert
from src.database import get_connection, DB_PATH

PROJECT_ROOT   = DB_PATH.parent.parent
RUTA_PLANTILLA = PROJECT_ROOT / "data" / "plantillas" / "plantilla_risst.docx"
CARPETA_FIRMAS = PROJECT_ROOT / "data" / "firmas"
CARPETA_SALIDA = PROJECT_ROOT / "documentos_generados"

_MESES = {1:\'enero\',2:\'febrero\',3:\'marzo\',4:\'abril\',5:\'mayo\',6:\'junio\',
          7:\'julio\',8:\'agosto\',9:\'septiembre\',10:\'octubre\',11:\'noviembre\',12:\'diciembre\'}


def convertir_fecha(fecha_str: str) -> str:
    try:
        f = datetime.strptime(fecha_str.strip(), \'%d/%m/%Y\')
        return f"Lima, {f.day} de {_MESES[f.month]} del {f.year}"
    except Exception:
        return fecha_str


def _aplicar_fuente(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.name=\'Century Gothic\'; r.font.size=Pt(11)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name=\'Century Gothic\'; r.font.size=Pt(11)


def generar_constancia(dni, nombre, apellido, cargo, fecha):
    try:
        CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)
        if not RUTA_PLANTILLA.exists():
            return False, f"Plantilla no encontrada: {RUTA_PLANTILLA}"
        doc = DocxTemplate(str(RUTA_PLANTILLA))
        ruta_firma = CARPETA_FIRMAS / f"firma_{dni}.png"
        contexto = {\'NOMBRE\':nombre, \'APELLIDOS\':apellido or nombre,
                    \'DNI\':dni, \'CARGO\':cargo or \'\', \'FECHA\':convertir_fecha(fecha)}
        if ruta_firma.exists():
            contexto[\'FIRMA\'] = InlineImage(doc, str(ruta_firma), width=Cm(4))
        doc.render(contexto)
        _aplicar_fuente(doc)
        nombre_docx = f"constancia_RISST_{nombre.replace(\' \',\'_\')}.docx"
        ruta_docx   = CARPETA_SALIDA / nombre_docx
        doc.save(str(ruta_docx))
        convert(str(ruta_docx))
        nombre_pdf = nombre_docx.replace(\'.docx\',\'.pdf\')
        with get_connection() as conn:
            conn.execute(
                "INSERT INTO documentos_generados (dni_trabajador,tipo_documento,fecha_generacion,ruta_archivo) VALUES (?,?,?,?)",
                (dni,\'CONSTANCIA_RISST\',datetime.now().strftime(\'%Y-%m-%d %H:%M:%S\'),str(CARPETA_SALIDA/nombre_pdf))
            )
        return True, nombre_pdf
    except Exception as exc:
        return False, str(exc)


def generar_todas_constancias(fecha: str) -> dict:
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT dni,nombre,apellido,cargo FROM trabajadores WHERE estado=\'ACTIVO\' ORDER BY nombre ASC"
        ).fetchall()
    trabajadores = [dict(r) for r in rows]
    exitosos=fallidos=0; log=[]
    for t in trabajadores:
        ok,msg = generar_constancia(t[\'dni\'],t[\'nombre\'],t[\'apellido\'] or \'\',t[\'cargo\'] or \'\',fecha)
        if ok: exitosos+=1; log.append(f"OK   {t[\'nombre\']}  ->  {msg}")
        else:  fallidos+=1; log.append(f"ERR  {t[\'nombre\']}  ->  {msg}")
    return {"total":len(trabajadores),"exitosos":exitosos,"fallidos":fallidos,"log":log}
'''

# ── src/procesar_firmas.py ───────────────────────────────────────────────────
files["src/procesar_firmas.py"] = '''\
from pathlib import Path
from PIL import Image
from rembg import remove
import io
from src.database import DB_PATH

PROJECT_ROOT    = DB_PATH.parent.parent
CARPETA_ENTRADA = PROJECT_ROOT / "data" / "firmas_originales"
CARPETA_SALIDA  = PROJECT_ROOT / "data" / "firmas"
ALTO_FINAL      = 65
_EXTENSIONES    = {\'.png\', \'.jpg\', \'.jpeg\'}


def procesar_firma(ruta_entrada: Path, ruta_salida: Path) -> None:
    with open(ruta_entrada, \'rb\') as f:
        datos = f.read()
    output = remove(datos)
    imagen = Image.open(io.BytesIO(output)).convert("RGBA")
    ancho_orig, alto_orig = imagen.size
    proporcion   = ALTO_FINAL / alto_orig
    ancho_final  = int(ancho_orig * proporcion)
    imagen_redim = imagen.resize((ancho_final, ALTO_FINAL), Image.Resampling.LANCZOS)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    imagen_redim.save(str(ruta_salida), format=\'PNG\')


def procesar_todas_firmas() -> dict:
    if not CARPETA_ENTRADA.exists():
        return {"procesadas":0,"omitidas":0,"errores":[f"Carpeta no encontrada: {CARPETA_ENTRADA}"]}
    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)
    procesadas=omitidas=0; errores=[]
    for archivo in sorted(CARPETA_ENTRADA.iterdir()):
        if archivo.suffix.lower() not in _EXTENSIONES:
            continue
        ruta_salida = CARPETA_SALIDA / (archivo.stem + \'.png\')
        if ruta_salida.exists():
            omitidas+=1; continue
        try:
            procesar_firma(archivo, ruta_salida)
            procesadas+=1
        except Exception as exc:
            errores.append(f"{archivo.name}: {exc}")
    return {"procesadas":procesadas,"omitidas":omitidas,"errores":errores}
'''

# ── src/gui.py ───────────────────────────────────────────────────────────────
files["src/gui.py"] = '''\
import threading
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk
from src.database import init_db
from src.generar_documento import generar_todas_constancias
from src.importar_excel import importar_trabajadores
from src.importar_sheets import sincronizar_desde_sheets
from src.procesar_firmas import procesar_todas_firmas
from src.trabajadores import buscar_por_dni, listar_activos


def _scrolled_text(parent, height):
    frame = ttk.Frame(parent)
    frame.pack(fill=\'both\', expand=True)
    widget = tk.Text(frame, height=height, font=(\'Consolas\',9), state=\'disabled\', wrap=\'none\')
    sb_y = ttk.Scrollbar(frame, orient=\'vertical\',   command=widget.yview)
    sb_x = ttk.Scrollbar(frame, orient=\'horizontal\', command=widget.xview)
    widget.configure(yscrollcommand=sb_y.set, xscrollcommand=sb_x.set)
    sb_y.pack(side=\'right\', fill=\'y\')
    sb_x.pack(side=\'bottom\', fill=\'x\')
    widget.pack(fill=\'both\', expand=True)
    return widget


def _write(widget, text, clear=True):
    widget.configure(state=\'normal\')
    if clear: widget.delete(\'1.0\',\'end\')
    widget.insert(\'end\', text)
    widget.see(\'end\')
    widget.configure(state=\'disabled\')


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("SGSST - MARA")
        self.geometry("720x580")
        self.resizable(False, False)
        s = ttk.Style()
        for w in (\'TLabel\',\'TButton\',\'TEntry\',\'TCombobox\'):
            s.configure(w, font=(\'Segoe UI\',10))
        nb = ttk.Notebook(self)
        nb.pack(fill=\'both\', expand=True, padx=10, pady=10)
        self._build_tab_trabajadores(nb)
        self._build_tab_constancias(nb)
        self._build_tab_firmas(nb)
        init_db()

    # ── TAB 1 ──────────────────────────────────────────────────────────────
    def _build_tab_trabajadores(self, nb):
        tab = ttk.Frame(nb, padding=15)
        nb.add(tab, text="  Trabajadores  ")
        acc = ttk.LabelFrame(tab, text="Importar / Sincronizar", padding=10)
        acc.pack(fill=\'x\', pady=(0,8))
        ttk.Button(acc, text="Sincronizar desde Google Sheets", command=self._sync_sheets).pack(side=\'left\', padx=(0,8))
        ttk.Button(acc, text="Importar desde Excel", command=self._importar_excel).pack(side=\'left\')
        srch = ttk.LabelFrame(tab, text="Consultar", padding=10)
        srch.pack(fill=\'x\', pady=(0,8))
        row = ttk.Frame(srch); row.pack(fill=\'x\')
        ttk.Label(row, text="DNI:").pack(side=\'left\', padx=(0,5))
        self._dni_var = tk.StringVar()
        ttk.Entry(row, textvariable=self._dni_var, width=14).pack(side=\'left\', padx=(0,8))
        ttk.Button(row, text="Buscar", command=self._buscar_dni).pack(side=\'left\', padx=(0,8))
        ttk.Button(row, text="Listar activos", command=self._listar_activos).pack(side=\'left\')
        res = ttk.LabelFrame(tab, text="Resultados", padding=10)
        res.pack(fill=\'both\', expand=True)
        self._txt_trabajadores = _scrolled_text(res, 14)

    def _sync_sheets(self):
        def run():
            try:
                res = sincronizar_desde_sheets()
                msg = (f"Sincronizacion completada\\n"
                       f"  Nuevos      : {res[\'nuevos\']}\\n"
                       f"  Actualizados: {res[\'actualizados\']}\\n"
                       f"  Omitidos    : {res[\'omitidos\']}")
                if res[\'errores\']: msg += "\\n\\nErrores:\\n" + "\\n".join(f"  {e}" for e in res[\'errores\'])
                self.after(0, lambda: messagebox.showinfo("Google Sheets", msg))
            except Exception as exc:
                self.after(0, lambda: messagebox.showerror("Error", str(exc)))
        threading.Thread(target=run, daemon=True).start()

    def _importar_excel(self):
        ruta = filedialog.askopenfilename(title="Seleccionar Excel", filetypes=[(\'Excel\',\'*.xlsx *.xls\'),(\'Todos\',\'*.*\')])
        if not ruta: return
        try:
            res = importar_trabajadores(ruta)
            msg = f"Importacion completada\\n  Importados: {res[\'importados\']}\\n  Omitidos  : {res[\'omitidos\']}"
            if res[\'errores\']: msg += "\\n\\nErrores:\\n" + "\\n".join(f"  {e}" for e in res[\'errores\'])
            messagebox.showinfo("Importar Excel", msg)
        except Exception as exc:
            messagebox.showerror("Error", str(exc))

    def _buscar_dni(self):
        dni = self._dni_var.get().strip()
        if not dni: messagebox.showwarning("Aviso","Ingresa un DNI"); return
        t = buscar_por_dni(dni)
        if t is None: _write(self._txt_trabajadores, "Trabajador no encontrado."); return
        _write(self._txt_trabajadores, "\\n".join(f"  {k:<20}: {v}" for k,v in t.items() if v is not None))

    def _listar_activos(self):
        activos = listar_activos()
        if not activos: _write(self._txt_trabajadores, "No hay trabajadores activos."); return
        wd = max(max(len(str(t[\'dni\'])) for t in activos),3)
        wn = max(max(len(str(t[\'nombre\'])) for t in activos),6)
        wc = max(max(len(str(t[\'cargo\'] or \'\')) for t in activos),5)
        sep = f"+-{\'-\'*wd}-+-{\'-\'*wn}-+-{\'-\'*wc}-+"
        lines = [sep, f"| {\'DNI\'.ljust(wd)} | {\'NOMBRE\'.ljust(wn)} | {\'CARGO\'.ljust(wc)} |", sep]
        for t in activos:
            lines.append(f"| {str(t[\'dni\']).ljust(wd)} | {str(t[\'nombre\']).ljust(wn)} | {str(t[\'cargo\'] or \'-\').ljust(wc)} |")
        lines += [sep, f"  Total: {len(activos)} trabajadores activos"]
        _write(self._txt_trabajadores, "\\n".join(lines))

    # ── TAB 2 ──────────────────────────────────────────────────────────────
    def _build_tab_constancias(self, nb):
        tab = ttk.Frame(nb, padding=15)
        nb.add(tab, text="  Constancias RISST  ")
        cfg = ttk.LabelFrame(tab, text="Configuracion", padding=10)
        cfg.pack(fill=\'x\', pady=(0,8))
        row1 = ttk.Frame(cfg); row1.pack(fill=\'x\', pady=3)
        ttk.Label(row1, text="Fecha (dd/mm/yyyy):").pack(side=\'left\', padx=(0,8))
        self._fecha_var = tk.StringVar(value=datetime.now().strftime(\'%d/%m/%Y\'))
        ttk.Entry(row1, textvariable=self._fecha_var, width=14).pack(side=\'left\')
        ttk.Label(cfg, text="Plantilla : data/plantillas/plantilla_risst.docx", foreground=\'gray\').pack(anchor=\'w\', pady=2)
        ttk.Label(cfg, text="Salida    : documentos_generados/", foreground=\'gray\').pack(anchor=\'w\')
        ttk.Button(tab, text="Generar constancias para todos los activos", command=self._generar_constancias).pack(pady=(0,6))
        self._prog_const = ttk.Progressbar(tab, mode=\'indeterminate\', length=500)
        self._prog_const.pack(fill=\'x\', pady=(0,8))
        log = ttk.LabelFrame(tab, text="Progreso", padding=10)
        log.pack(fill=\'both\', expand=True)
        self._txt_constancias = _scrolled_text(log, 12)

    def _generar_constancias(self):
        fecha = self._fecha_var.get().strip()
        if not fecha: messagebox.showwarning("Aviso","Ingresa una fecha"); return
        _write(self._txt_constancias, "")
        self._prog_const.start()
        def run():
            try:
                res = generar_todas_constancias(fecha)
                for line in res[\'log\']:
                    self.after(0, lambda l=line: _write(self._txt_constancias, l+\'\\n\', clear=False))
                summary = f"\\n{\'=\'*50}\\nTotal:{res[\'total\']} | OK:{res[\'exitosos\']} | Errores:{res[\'fallidos\']}"
                self.after(0, lambda: _write(self._txt_constancias, summary, clear=False))
                self.after(0, lambda: messagebox.showinfo("Completado",
                    f"Generacion finalizada.\\n{res[\'exitosos\']} de {res[\'total\']} constancias generadas."))
            except Exception as exc:
                self.after(0, lambda: messagebox.showerror("Error", str(exc)))
            finally:
                self.after(0, self._prog_const.stop)
        threading.Thread(target=run, daemon=True).start()

    # ── TAB 3 ──────────────────────────────────────────────────────────────
    def _build_tab_firmas(self, nb):
        tab = ttk.Frame(nb, padding=15)
        nb.add(tab, text="  Firmas  ")
        info = ttk.LabelFrame(tab, text="Instrucciones", padding=10)
        info.pack(fill=\'x\', pady=(0,10))
        ttk.Label(info, text=(
            "1. Coloca las imagenes en:  data/firmas_originales/\\n"
            "2. Nombre esperado: firma_{DNI}.png  (ej. firma_05314709.png)\\n"
            "3. Presiona \'Procesar firmas\'.\\n"
            "4. Las imagenes sin fondo quedaran en: data/firmas/\\n\\n"
            "Archivos ya procesados se omiten automaticamente."
        ), justify=\'left\', foreground=\'#333333\').pack(anchor=\'w\')
        ttk.Button(tab, text="Procesar firmas", command=self._procesar_firmas).pack(pady=(0,6))
        self._prog_firmas = ttk.Progressbar(tab, mode=\'indeterminate\', length=500)
        self._prog_firmas.pack(fill=\'x\', pady=(0,8))
        log = ttk.LabelFrame(tab, text="Progreso", padding=10)
        log.pack(fill=\'both\', expand=True)
        self._txt_firmas = _scrolled_text(log, 14)

    def _procesar_firmas(self):
        _write(self._txt_firmas, "")
        self._prog_firmas.start()
        def run():
            try:
                res = procesar_todas_firmas()
                lines = [f"Procesadas : {res[\'procesadas\']}", f"Omitidas   : {res[\'omitidas\']}  (ya existian)"]
                for err in res[\'errores\']: lines.append(f"Error      : {err}")
                self.after(0, lambda: _write(self._txt_firmas, "\\n".join(lines)))
                self.after(0, lambda: messagebox.showinfo("Completado",
                    f"Procesadas: {res[\'procesadas\']}\\nOmitidas: {res[\'omitidas\']}\\nErrores: {len(res[\'errores\'])}"))
            except Exception as exc:
                self.after(0, lambda: messagebox.showerror("Error", str(exc)))
            finally:
                self.after(0, self._prog_firmas.stop)
        threading.Thread(target=run, daemon=True).start()


def main():
    App().mainloop()


if __name__ == "__main__":
    main()
'''

# ── Write all files ──────────────────────────────────────────────────────────
for rel_path, content in files.items():
    dest = BASE / rel_path
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(content, encoding="utf-8")
    print(f"  OK  {rel_path}")

print("\nTodos los archivos creados. Ahora ejecuta:")
print("  pip install -r requirements.txt")
print("  python -m src.gui")
